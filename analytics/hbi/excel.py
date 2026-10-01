"""Write a schedule to a colour-coded Excel workbook in the layout managers already use."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schedule import SHIFTS

FILL = {
    "AM": "DCEBFB", "MID": "DFF3E4", "SH": "FFF4CC", "INS": "E9E2F8", "PM": "FDE3CF",
    "OFF": "F2F2F2", "N/A": "D9D9D9",
}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_schedule(path, title: str, staff: list, grid: dict, demand, coverage: list, rules: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    days = [f"{r['day']} {r['date'][5:]}" for _, r in demand.iterrows()]

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Codes: AM 07:00-15:30, MID 09:00-17:30, SH 09:00-13:00, INS inspector 07:00-15:30, PM 15:00-23:00"
    ws["A2"].font = Font(italic=True, size=9, color="595959")

    header = ["Employee", "Name", "Contract"] + days + ["Hours"]
    for c, h in enumerate(header, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = BOX

    hours_ws = wb.create_sheet("Hours")
    hours_ws.append(["Employee"] + days)
    for i, e in enumerate(staff, start=5):
        g = grid[e.emp_id]
        ws.cell(row=i, column=1, value=e.emp_id).border = BOX
        ws.cell(row=i, column=2, value=e.name).border = BOX
        ws.cell(row=i, column=3, value=e.contract).border = BOX
        hours_ws.append([e.emp_id] + [SHIFTS[s][4] if s in SHIFTS else 0 for s in g])
        for d, s in enumerate(g):
            cell = ws.cell(row=i, column=4 + d, value=s)
            cell.fill = PatternFill("solid", fgColor=FILL.get(s, "FFFFFF"))
            cell.alignment = Alignment(horizontal="center")
            cell.border = BOX
        hrow = i - 3
        total = ws.cell(row=i, column=11, value=f"=SUM(Hours!B{hrow}:H{hrow})")
        total.border = BOX
        total.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 9
    for c in range(4, 12):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.freeze_panes = "D5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    cv = wb.create_sheet("Coverage")
    cols = list(coverage[0].keys())
    cv.append([c.replace("_", " ").capitalize() for c in cols])
    for r in coverage:
        cv.append([r[c] for c in cols])
    for c in range(1, len(cols) + 1):
        cv.cell(row=1, column=c).font = Font(bold=True)
        cv.column_dimensions[get_column_letter(c)].width = 22

    rc = wb.create_sheet("Rule check")
    rc.append(["Rule", "Violations", "Result"])
    for r in rules:
        rc.append([r["rule"], r["violations"], "Pass" if r["passed"] else "Fail"])
    rc.column_dimensions["A"].width = 90
    for c in range(1, 4):
        rc.cell(row=1, column=c).font = Font(bold=True)

    dm = wb.create_sheet("Demand")
    dm.append([c.replace("_", " ").capitalize() for c in demand.columns])
    for _, r in demand.iterrows():
        dm.append(list(r.values))
    for c in range(1, len(demand.columns) + 1):
        dm.cell(row=1, column=c).font = Font(bold=True)
        dm.column_dimensions[get_column_letter(c)].width = 18
    wb.save(path)
